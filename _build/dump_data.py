# -*- coding: utf-8 -*-
"""读取 13180操作系统.xlsx（13180操作系统大纲 sheet）全量数据并落盘 JSON。"""
import json
from openpyxl import load_workbook

SRC = r"F:\syncthing\考试\自考\13180操作系统\13180操作系统.xlsx"
OUT = r"F:\syncthing\考试\自考\13180操作系统\_build\rows.json"

wb = load_workbook(SRC, data_only=True, read_only=True)
ws = wb["13180操作系统大纲"]

rows = []
for i, row in enumerate(ws.iter_rows(min_row=1, max_col=7, values_only=True), start=1):
    rows.append({
        "excel_row": i,
        "chapter": (row[0] if row[0] is not None else ""),
        "requirement": (row[1] if row[1] is not None else ""),
        "summary": (row[2] if row[2] is not None else ""),
        "knowledge": (row[3] if row[3] is not None else ""),
        "desc": (row[4] if row[4] is not None else ""),
        "seq": row[5],
        "page": row[6],
    })
wb.close()

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(rows, f, ensure_ascii=False, indent=1)

total = len(rows)
has_k = sum(1 for r in rows if r["knowledge"].strip())
has_e = sum(1 for r in rows if r["desc"].strip())
has_page = sum(1 for r in rows if r["page"] is not None)
chapters, reqs, summaries = [], [], []
for r in rows:
    ch = r["chapter"].strip()
    if ch and ch not in chapters:
        chapters.append(ch)
    q = r["requirement"].strip()
    if q and q not in reqs:
        reqs.append(q)
    s = r["summary"].strip()
    if s and s not in summaries:
        summaries.append(s)
lens = [len(r["desc"]) for r in rows if r["desc"].strip()]
print("总行数:", total)
print("含知识点D:", has_k, "含描述E:", has_e, "含页码G:", has_page)
print("章节数:", len(chapters), chapters)
print("要求B取值:", reqs)
print("小结C数量:", len(summaries))
print("E长度: min=%d max=%d avg=%.1f" % (min(lens), max(lens), sum(lens) / len(lens)))
print("小结C取值:", summaries)
